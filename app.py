import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
import io
import json
import google.generativeai as genai

# 1. Cấu hình giao diện trang web
st.set_page_config(
    page_title="AI Lesson Plan to PPTX", 
    page_icon="📊", 
    layout="wide"
)

st.title("🎓 Trợ Lý AI Chuyển Kế Hoạch Bài Dạy Thành PowerPoint (.pptx)")
st.write("Hỗ trợ giáo viên phân tích Kế hoạch bài dạy và tự động tạo tệp slide trình chiếu (.pptx) phục vụ giảng dạy.")

# 2. Thanh cấu hình API Key ở Sidebar
with st.sidebar:
    st.header("⚙️ Cấu hình API")
    api_key = st.text_input("Nhập Gemini API Key của thầy/cô:", type="password")
    st.markdown("[👉 Lấy Gemini API Key miễn phí tại đây](https://aistudio.google.com/app/apikey)")

# 3. Form nhập nội dung kế hoạch bài dạy
st.subheader("📝 Nhập Kế hoạch bài dạy (Giáo án)")
lesson_plan = st.text_area(
    "Dán nội dung Kế hoạch bài dạy/Giáo án vào đây:",
    height=250,
    placeholder="Nhập hoặc dán nội dung kế hoạch bài dạy (Mục tiêu, Tiến trình dạy học, các hoạt động...)"
)

# 4. Xử lý tạo Slide PowerPoint khi bấm nút
if st.button("🚀 Khởi Tạo Slide PowerPoint (.pptx)", type="primary"):
    if not api_key:
        st.error("⚠️ Vui lòng nhập Gemini API Key ở thanh bên trái!")
    elif not lesson_plan.strip():
        st.warning("⚠️ Vui lòng dán nội dung Kế hoạch bài dạy!")
    else:
        try:
            with st.spinner("⏳ AI đang phân tích bài dạy và thiết kế bài thuyết trình..."):
                # Cấu hình Gemini API
                genai.configure(api_key=api_key)
                model = genai.GenerativeModel('gemini-2.0-flash')
                
                # System Prompt yêu cầu AI xuất ra cấu trúc JSON
                prompt_yeu_cau = f"""
Bạn là một chuyên gia thiết kế bài giảng điện tử.
Hãy phân tích Kế hoạch bài dạy dưới đây và chuyển thành danh sách các Slide trình chiếu phục vụ giảng dạy.
BẮT BUỘC trả về kết quả dưới dạng cấu trúc JSON hợp lệ, không chứa ký tự markdown thừa, đúng định dạng sau:
{{
  "presentation_title": "Tên bài học",
  "slides": [
    {{
      "title": "Tiêu đề Slide",
      "content": [
        "Ý thứ nhất ngắn gọn, xúc tích",
        "Ý thứ hai ngắn gọn, xúc tích",
        "Ý thứ ba ngắn gọn, xúc tích"
      ]
    }}
  ]
}}

NỘI DUNG KẾ HOẠCH BÀI DẠY:
{lesson_plan}
"""
                response = model.generate_content(prompt_yeu_cau)
                
                # Làm sạch chuỗi JSON từ phản hồi của Gemini
                json_text = response.text.replace("```json", "").replace("```", "").strip()
                data = json.loads(json_text)
                
                # Khởi tạo tệp PowerPoint bằng python-pptx
                prs = Presentation()
                
                # Slide 1: Slide Tiêu đề
                title_layout = prs.slide_layouts
                slide_title = prs.slides.add_slide(title_layout)
                slide_title.shapes.title.text = data.get("presentation_title", "Bài Giảng Điện Tử")
                slide_title.placeholders[10].text = "Slide bài giảng được tạo tự động bởi Trợ lý AI"
                
                # Các Slide nội dung tiếp theo
                bullet_layout = prs.slide_layouts[10]
                for item in data.get("slides", []):
                    slide = prs.slides.add_slide(bullet_layout)
                    slide.shapes.title.text = item.get("title", "Nội dung bài học")
                    
                    text_frame = slide.placeholders[10].text_frame
                    text_frame.word_wrap = True
                    
                    contents = item.get("content", [])
                    if contents:
                        text_frame.text = contents
                        for line in contents[1:]:
                            p = text_frame.add_paragraph()
                            p.text = line
                
                # Luồng lưu file nhị phân vào bộ nhớ io.BytesIO
                binary_output = io.BytesIO()
                prs.save(binary_output)
                binary_output.seek(0)
                
                st.success("🎉 Tạo bộ slide PowerPoint thành công!")
                
                # Nút tải xuống file PowerPoint (.pptx)
                st.download_button(
                    label="📥 Tải Xuống Tệp PowerPoint (.pptx)",
                    data=binary_output.getvalue(),
                    file_name=f"Bai_giang_{data.get('presentation_title', 'slide').replace(' ', '_')}.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
                )
        except Exception as e:
            st.error(f"❌ Có lỗi xảy ra trong quá trình xử lý: {str(e)}")